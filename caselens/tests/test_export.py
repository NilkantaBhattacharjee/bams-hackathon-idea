from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import export


def _frames():
    by_type = pd.DataFrame({"Type": ["Retail_Banking", "Commercial_Lending"], "Cases": [3, 2],
                            "Stalled": [1, 0], "Avg_TAT_hours": [40.0, 55.5], "Stalled_%": [33.3, 0.0]})
    counts = pd.DataFrame({"Stage": [2, 3], "Cases": [2, 1], "Stage_name": ["Verification", "Assessment"],
                           "Stage_label": ["S2 Verification", "S3 Assessment"]})
    sub = pd.DataFrame({"Sub_type": ["Personal_Loan", "SME_Working_Capital"], "Cases": [3, 2],
                        "Stalled": [1, 0], "Stalled_%": [33.3, 0.0]})
    cases = pd.DataFrame({"Case_ID": ["CS-2026-1001", "CS-2026-1002"], "Created": ["01-Mar-2026", "02-Mar-2026"],
                          "Type": ["Retail_Banking", "Commercial_Lending"],
                          "Sub_type": ["Personal_Loan", "SME_Working_Capital"],
                          "TAT_hours": [40.0, None], "Stalled": [True, False], "Closed": [True, False]})
    report = [("Date range", "01 Mar 2026 to 02 Mar 2026"), ("Cases", "2"), ("Trend report", "[T1] text")]
    return report, by_type, counts, sub, cases


def test_workbook_has_all_sheets_and_charts():
    data = export.build_workbook(*_frames())
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Report", "By type", "Longest stage", "Sub-type", "Cases"]
    assert len(wb["By type"]._charts) == 2
    assert len(wb["Longest stage"]._charts) == 1
    assert len(wb["Sub-type"]._charts) == 1
    assert wb["Report"]["B2"].value == "01 Mar 2026 to 02 Mar 2026"
    assert wb["Cases"]["A2"].value == "CS-2026-1001"


def test_empty_tables_do_not_break_export():
    report, by_type, counts, sub, cases = _frames()
    empty = by_type.iloc[0:0]
    data = export.build_workbook(report, empty, counts.iloc[0:0], sub.iloc[0:0], cases.iloc[0:0])
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames[0] == "Report"
