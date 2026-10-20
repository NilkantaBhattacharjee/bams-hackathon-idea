"""Excel export for the bulk view. One workbook holds the trend report, every table and the charts.

Each sheet is built from the same DataFrames the page shows, so the export matches the screen.
"""

from io import BytesIO

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="6C5CE7")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_sheet(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 12), 60)


def _bar_chart(ws, n_rows: int, cat_col: str, value_cols: list[str], title: str, y_title: str, anchor: str) -> None:
    """Native Excel bar chart built from the sheet's own table (header in row 1, data below)."""
    headers = [c.value for c in ws[1]]
    chart = BarChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.width, chart.height = 16, 8
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    for name in value_cols:
        col = headers.index(name) + 1
        chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=n_rows + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=headers.index(cat_col) + 1, min_row=2, max_row=n_rows + 1))
    ws.add_chart(chart, anchor)


def build_workbook(report_rows, by_type: pd.DataFrame, counts: pd.DataFrame,
                   sub: pd.DataFrame, cases: pd.DataFrame) -> bytes:
    """Return the .xlsx file as bytes.

    report_rows: list of (section, detail) pairs for the Report sheet.
    by_type, counts, sub, cases: the tables shown on the bulk page.
    """
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(report_rows, columns=["Section", "Detail"]).to_excel(xw, sheet_name="Report", index=False)
        by_type.to_excel(xw, sheet_name="By type", index=False)
        counts[["Stage", "Stage_name", "Cases"]].to_excel(xw, sheet_name="Longest stage", index=False)
        sub.to_excel(xw, sheet_name="Sub-type", index=False)
        cases.to_excel(xw, sheet_name="Cases", index=False)

        for ws in xw.book.worksheets:
            _style_sheet(ws)
        for cell in xw.sheets["Report"]["B"][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        if len(by_type):
            ws = xw.sheets["By type"]
            _bar_chart(ws, len(by_type), "Type", ["Cases", "Stalled"],
                       "Cases and stalled cases by type", "Cases", "H2")
            _bar_chart(ws, len(by_type), "Type", ["Avg_TAT_hours"],
                       "Average turnaround by type", "Hours", "H20")
        if len(counts):
            _bar_chart(xw.sheets["Longest stage"], len(counts), "Stage_name", ["Cases"],
                       "Longest stage per case", "Cases", "E2")
        if len(sub):
            _bar_chart(xw.sheets["Sub-type"], len(sub), "Sub_type", ["Stalled_%"],
                       "Stalled cases (%) by sub-type", "Percent", "F2")
    return buf.getvalue()
