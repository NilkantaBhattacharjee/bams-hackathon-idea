"""Builds the 5-slide CaseLens deck: docs/deliverables/CaseLens_Deck.pptx

Run from the caselens folder:  py ..\\docs\\deliverables\\make_deck.py
Needs:  py -m pip install python-pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "CaseLens_Deck.pptx"
PURPLE = RGBColor(0x6C, 0x5C, 0xE7)
DARK = RGBColor(0x2D, 0x34, 0x36)
GREY = RGBColor(0x63, 0x6E, 0x72)


def add_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = title
    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run.font.size = Pt(30)
    title_run.font.bold = True
    title_run.font.color.rgb = PURPLE
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.4))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(bullets):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        is_heading = line.startswith("#")
        para.text = line.lstrip("# ").strip() if is_heading else "- " + line
        para.font.size = Pt(18 if is_heading else 16)
        para.font.bold = is_heading
        para.font.color.rgb = DARK if is_heading else GREY
        para.space_after = Pt(6)
    return slide


def add_table_slide(prs, title, rows, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run.font.size = Pt(30)
    title_run.font.bold = True
    title_run.font.color.rgb = PURPLE
    n_rows, n_cols = len(rows), len(rows[0])
    table = slide.shapes.add_table(n_rows, n_cols, Inches(0.6), Inches(1.6),
                                   Inches(12.1), Inches(0.45 * n_rows)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(14)
                para.font.bold = r == 0
    if note:
        box = slide.shapes.add_textbox(Inches(0.6), Inches(6.4), Inches(12.1), Inches(0.8))
        box.text_frame.word_wrap = True
        box.text_frame.text = note
        box.text_frame.paragraphs[0].font.size = Pt(13)
        box.text_frame.paragraphs[0].font.color.rgb = GREY
    return slide


def add_image_slide(prs, title, image_path, caption):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run.font.size = Pt(30)
    title_run.font.bold = True
    title_run.font.color.rgb = PURPLE
    if Path(image_path).exists():
        slide.shapes.add_picture(str(image_path), Inches(0.8), Inches(1.5), height=Inches(5.0))
    box = slide.shapes.add_textbox(Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.7))
    box.text_frame.word_wrap = True
    box.text_frame.text = caption
    box.text_frame.paragraphs[0].font.size = Pt(13)
    box.text_frame.paragraphs[0].font.color.rgb = GREY
    return slide


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # Slide 1: persona, pain, success metrics
    add_slide(prs, "CaseLens: Persona, Pain and Success Metrics", [
        "# Persona",
        "Priya, shift lead in retail banking operations. Manages about 40 BPM cases a shift.",
        "# Pain",
        "Case comments are abbreviated (KYC, PEP, LTV, DTI, CBS) across five stages.",
        "Timed trial: 20 cases took about 20 minutes by hand (about 1 minute per case).",
        "Managers need open and stalled counts with average turnaround for a period.",
        "Auditors need trends and charts, not case-by-case reading.",
        "# Success metrics",
        "Review time: 7 to 8 minutes with CaseLens vs 20 minutes by hand (60 to 65% faster).",
        "Grounding: 100% of summary lines cite a source (5,808 of 5,808).",
        "Invented terms refused: 30 of 30. Raw client IDs in output: 0.",
        "p95 latency: 2.5 seconds per model call.",
    ])

    # Slide 2: architecture diagram
    add_image_slide(prs, "Architecture",
                    Path(__file__).resolve().parent / "architecture.png",
                    "Facts are computed in code. The model only explains them, and every line must cite a source. "
                    "Source: 05_architecture.md. Export the Mermaid diagram as PNG from mermaid.live and save it as "
                    "architecture.png in this folder.")

    # Slide 3: evaluation evidence
    add_table_slide(prs, "Evaluation Evidence", [
        ["Check", "Target", "Result"],
        ["Grounding (cited summary lines)", "95% or more", "100% (5,808 of 5,808)"],
        ["Invented terms refused", "100%", "100% (30 of 30)"],
        ["Raw client IDs in output", "0", "0"],
        ["p95 latency per model call", "under 5 s", "2.5 s"],
        ["Live evaluation: summaries", "pass", "10 of 10"],
        ["Live evaluation: chat", "pass", "9 of 9"],
        ["Automated test suite", "all pass", "150 of 150"],
    ], note="Failure cases fixed: a valid question refused as 'not in glossary'; a time-to-stage answer printing 'nan'. "
            "Open: retrieval not measured on real staff questions; feedback not scored.")

    # Slide 4: risk and control
    add_table_slide(prs, "Risk and Control Notes", [
        ["Risk", "Control in place", "Residual"],
        ["Model invents a definition", "Glossary-only answers; 30-term refusal check", "Low"],
        ["Client data sent to model", "Masking before every call; personal-data questions refused", "Low"],
        ["Prompt injection", "Override-phrase scan; comments treated as data", "Low to medium"],
        ["Uncited claim reaches user", "Citation check drops and counts uncited lines", "Low"],
        ["API key exposed", "Key in local .env, never logged; key rotated 22 Oct", "Low"],
        ["Users over-trust answers", "Decision-support label; thumbs up/down; human decides", "Medium"],
    ], note="Compliance: all data is synthetic. A real deployment needs a data-protection review, "
            "including lawful basis, log retention, and provider data terms.")

    # Slide 5: cost and scalability
    add_table_slide(prs, "Cost Estimate and Scalability", [
        ["Item", "Estimate"],
        ["Model", "openai/gpt-4o-mini via OpenRouter"],
        ["Price used", "$0.15 per M input, $0.60 per M output (blended $0.000375 per 1K tokens)"],
        ["Tokens per summary", "about 1,400"],
        ["Cost per summary", "about $0.0005"],
        ["Cost per 1,000 summaries", "about $0.53 (model only)"],
        ["Infrastructure", "Not yet estimated (prototype runs locally)"],
    ], note="Scaling: cost grows with summaries requested, not cases stored. Summary cache removes repeat cost. "
            "Data moves to the bank's BPM database via read-only connection. Prices must be re-checked before use.")

    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
